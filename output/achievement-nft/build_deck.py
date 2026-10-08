"""Judge appendix: editable diagrams + matching PNG/PDF. No wallet transactions."""
from pathlib import Path
from html import escape
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

OUT=Path(__file__).resolve().parent
(OUT/'slides').mkdir(exist_ok=True)
prs=Presentation(); prs.slide_width=Inches(13.333333); prs.slide_height=Inches(7.5)
images=[]; notes=[]
BG='070B16'; PANEL='111A2B'; WHITE='F3F7FF'; MUTED='A8B6CF'; MINT='35E5C5'; GOLD='FFD080'
def font(size):
    return ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial.ttf',size)
def box(x,y,w,h,c=PANEL):
    q=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x/144),Inches(y/144),Inches(w/144),Inches(h/144))
    q.fill.solid();q.fill.fore_color.rgb=RGBColor.from_string(c);q.line.fill.background()
    draw.rounded_rectangle((x,y,x+w,y+h),radius=16,fill='#'+c)
def text(value,x,y,w,size=32,c=WHITE):
    lines=value.split('\n'); h=len(lines)*size*1.3+12
    assert x+w<=1920 and y+h<=1080
    for line in lines: assert draw.textlength(line,font=font(size))<=w,(line,w)
    q=s.shapes.add_textbox(Inches(x/144),Inches(y/144),Inches(w/144),Inches(h/144))
    tf=q.text_frame;tf.word_wrap=False;tf.margin_left=tf.margin_top=tf.margin_bottom=tf.margin_right=0
    for i,line in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.space_after=Pt(0);p.line_spacing=Pt(size*1.3/2)
        r=p.add_run();r.text=line;r.font.name='Arial';r.font.size=Pt(size/2);r.font.color.rgb=RGBColor.from_string(c)
        draw.text((x,y+i*size*1.3),line,font=font(size),fill='#'+c)
def start(n,title):
    global s,canvas,draw
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(BG)
    canvas=Image.new('RGB',(1920,1080),'#'+BG);draw=ImageDraw.Draw(canvas)
    text('NEONSHIFT / PERSONAL ACHIEVEMENT NFT',80,55,1760,25,MINT)
    text(title,80,120,1760,62)
    text(f'JUDGE APPENDIX {n}/3  /  DESIGN PREVIEW - NOT A MINT RECEIPT',80,1010,1760,23,MUTED)
def finish(n,note):
    s.notes_slide.notes_text_frame.text=note
    notes.append(note);images.append(canvas.copy());canvas.save(OUT/'slides'/f'{n:02}.png')

start(1,'A personal record, a collectible identity.')
cards=[('5','FIRST DISTANCE','35E5C5'),('10','FIRST DISTANCE','62CFFF'),('21.0975','HALF MARATHON','AD8CFF'),('42.195','MARATHON','FFD080')]
for i,(distance,label,c) in enumerate(cards):
    x=80+i*448;box(x,260,416,490)
    text('GENESIS DISTANCE',x+25,290,370,25,c)
    box(x+30,355,356,7,c)
    text(distance,x+25,405,370,68,c)
    text('km / '+label,x+25,520,370,22)
    text('DEVICE RECORDED',x+25,603,370,22,MUTED)
    text('SAMPLE / NO OWNER DATA',x+25,685,370,18,MUTED)
text('VISUAL SYSTEM',80,802,1760,28,MINT)
text('Distance first. A distinct color per milestone. Source labels stay visible.',80,851,1760,36)
text('Other series: PB / First Finish / Event Memory. Organizer and device records remain distinct.',80,920,1760,29,MUTED)
finish(1,'中文講稿：這是個人成就 NFT 的資訊版式提案，不是已鑄造證明。距離是主角，來源標籤不能省略。既有正式 SVG 圖樣請看同目錄 preview.html。首次距離與 First Finish 可從 Lv1 開始；PB 仍依達成時 Lv3 等資格規則。裝置來源上線狀態以 App 實際資格頁為準。')

start(2,'From verified achievement to wallet ownership.')
steps=[('01 / RECORD','Complete a qualifying run.\nSync it to NeonShift.'),('02 / VERIFY','Check source and eligibility.\nWait for registry approval.'),('03 / PREVIEW','Choose public details.\nReview metadata and SOL fee.'),('04 / APPROVE','Approve in your wallet.\nWait for chain confirmation.')]
for i,(title,body) in enumerate(steps):
    x=80+i*448;box(x,280,416,310);text(title,x+22,310,372,30,MINT);text(body,x+22,395,372,25)
text('CLAIM = USER-INITIATED MINT',80,650,1760,33,GOLD)
text('Reaching a target does not automatically mint an NFT.',80,710,1760,41)
text('No tSKR spend. Review the current SOL estimate before approving.',80,780,1760,34)
text('Public blockchain ownership remains visible; optional detail sharing is not anonymity.',80,860,1760,31,MUTED)
text('Raw GPS, heart rate and weight are not included in the achievement NFT.',80,915,1760,31,MUTED)
finish(2,'中文講稿：完成運動只是開始。系統驗證來源、門檻和唯一性，registry 核准後才可鑄造。使用者選擇公開內容、查看 metadata 與當次 SOL 費用，再到錢包核准。首次距離的領取就是主動發起鑄造，不是先自動鑄造後再空投。依據 Milestones.tsx 與 AchievementService.ts；此頁是流程圖，尚未附上真實交易。')

start(3,'Claim once. Keep the story. Show the evidence.')
for y,title,body in [
(260,'ELIGIBLE IS NOT MINTED','Locked / under review / pending approval stay separate from claimed.'),
(430,'CONFIRMED AND COLLECTED','Show the reveal, then open the collection and inspect the asset.'),
(600,'RETRIES DO NOT CREATE DUPLICATES','An existing on-chain receipt returns the existing asset instead of minting again.')]:
    box(80,y,1760,145);text(title,110,y+22,1690,29,MINT);text(body,110,y+77,1690,31)
text('LIVE DEMO CHECKLIST',80,805,1760,28,GOLD)
text('Eligibility -> privacy + fee preview -> wallet approval -> confirmation -> collection',80,860,1760,31)
text('Show the actual asset and transaction on Devnet. Never replace them with sample IDs.',80,925,1760,29,MUTED)
finish(3,'中文講稿：評審應看到可領取、待核准、已領取是不同狀態。拒簽不發成功动画；重試先查 receipt，避免重複鑄造。實錄需補上提交版本的資格畫面、公開內容與費用、錢包核准、成功收藏及真實 Devnet asset／交易。此附件沒有執行鑄造，也沒有捏造交易網址。')
prs.save(OUT/'NeonShift_Achievement_NFT_Judges_EN.pptx')
images[0].save(OUT/'NeonShift_Achievement_NFT_Judges_EN.pdf',save_all=True,append_images=images[1:],resolution=144)
(OUT/'講稿.md').write_text('# 評審 NFT 展示講稿\n\n'+'\n\n'.join(f'## {i+1}\n\n{n}' for i,n in enumerate(notes))+'\n')
assets=['first_5k-device','first_10k-device','first_half-device','first_marathon-device','first_5k-organizer','first_finish-organizer']
(OUT/'preview.html').write_text('<!doctype html><html lang="zh-Hant"><meta charset="utf-8"><title>NeonShift NFT 評審展示</title><style>body{background:#070b16;color:#f3f7ff;font:18px system-ui;margin:32px auto;max-width:1200px}img.slide{width:100%}.gallery{display:grid;grid-template-columns:repeat(auto-fit,minmax(280px,1fr));gap:20px}.gallery img{width:100%}p{line-height:1.7}a{color:#35e5c5}</style><h1>個人成就 NFT：樣式與領取流程</h1><p>投影片是設計／流程示意，不是鑄造紀錄。下方提供專案既有 SVG 圖樣；來源分類不等於目前已開放領取。</p>'+''.join(f'<img class="slide" alt="評審展示第 {i} 頁" src="slides/{i:02}.png">' for i in range(1,4))+'<h2>現有 NFT 圖樣</h2><div class="gallery">'+''.join(f'<figure><img alt="{escape(a)}" src="../../web/nft/achievements/milestones/{a}.svg"><figcaption>{escape(a)}</figcaption></figure>' for a in assets)+'</div><p>資訊版式提案不覆寫現有 NFT metadata 或已發行資產。裝置版是否可領取以 App 資格狀態為準。</p></html>')
print('Built three editable judge slides, PDF, PNG, HTML gallery and Chinese notes.')
# Review edition: preserve the six-scene video deck, append three editable Q&A slides.
from copy import deepcopy
main=OUT.parent/'hackathon-flexclip-en'
combined=Presentation(main/'NeonShift_Hackathon_EN.pptx')
for source in prs.slides:
    dest=combined.slides.add_slide(combined.slide_layouts[6])
    dest.background.fill.solid();dest.background.fill.fore_color.rgb=RGBColor.from_string(BG)
    for shape in source.shapes:
        dest.shapes._spTree.insert_element_before(deepcopy(shape.element),'p:extLst')
    dest.notes_slide.notes_text_frame.text=source.notes_slide.notes_text_frame.text
combined.save(main/'NeonShift_Hackathon_EN_Judges.pptx')
overview=Image.new('RGB',(960,1620),'#'+BG)
for i,im in enumerate(images): overview.paste(im.resize((960,540)),(0,i*540))
overview.save(OUT/'overview.jpg')
