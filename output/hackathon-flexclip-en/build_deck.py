"""Build editable PPTX and matching 1920x1080 composition exports. Run with python-pptx and Pillow."""
from pathlib import Path
from html import escape
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
import json

OUT = Path(__file__).resolve().parent
BG='070B16'; PANEL='111A2B'; WHITE='F3F7FF'; MUTED='A8B6CF'; MINT='35E5C5'; PURPLE='AD8CFF'; GOLD='FFD080'
FONT_DIR=Path('/System/Library/Fonts/Supplemental')
prs=Presentation(); prs.slide_width=Inches(13.333333); prs.slide_height=Inches(7.5)
slides=[]; images=[]; manifest=[]
NARRATION=(OUT/'voiceover-v3-archive.txt').read_text().strip().split('\n\n')
assert len(NARRATION)==6

def font(size,bold=False):
    path=FONT_DIR/('Arial Bold.ttf' if bold else 'Arial.ttf')
    if not path.exists():
        path=Path('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf' if bold else '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')
    return ImageFont.truetype(str(path),size)

def box(x,y,w,h,color=PANEL):
    shape=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x/144), Inches(y/144), Inches(w/144), Inches(h/144))
    shape.fill.solid(); shape.fill.fore_color.rgb=RGBColor.from_string(color); shape.line.fill.background()
    draw.rectangle((x,y,x+w,y+h),fill='#'+color)

def txt(t,x,y,w,size=36,color=WHITE,bold=False):
    lines=t.split('\n'); lh=round(size*1.23); h=lh*len(lines)+8
    assert x+w<=1920 and y+h<=1080, (t,x,y,w,h)
    tb=s.shapes.add_textbox(Inches(x/144),Inches(y/144),Inches(w/144), Inches(h/144))
    tf=tb.text_frame; tf.word_wrap=False; tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(lines):
        assert draw.textlength(line,font=font(size,bold)) <= w, ('text too wide',line,w)
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.space_after=Pt(0); p.space_before=Pt(0); p.line_spacing=Pt(lh/2)
        r=p.add_run(); r.text=line; r.font.name='Arial'; r.font.size=Pt(size/2); r.font.bold=bold; r.font.color.rgb=RGBColor.from_string(color)
        draw.text((x,y+i*lh),line,font=font(size,bold),fill='#'+color,anchor='lt')

def pic(name,x,y,w,h):
    path=OUT/'assets'/name; im=Image.open(path).convert('RGB'); ratio=min(w/im.width,h/im.height)
    iw,ih=round(im.width*ratio),round(im.height*ratio); xx=x+(w-iw)//2; yy=y+(h-ih)//2
    canvas.paste(im.resize((iw,ih),Image.Resampling.LANCZOS),(xx,yy))
    s.shapes.add_picture(str(path),Inches(xx/144),Inches(yy/144),width=Inches(iw/144),height=Inches(ih/144))

def start(n,kicker):
    global s,canvas,draw
    s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=RGBColor.from_string(BG)
    canvas=Image.new('RGB',(1920,1080),'#'+BG); draw=ImageDraw.Draw(canvas)
    box(80,64,10,24,MINT); txt('NEONSHIFT  /  '+kicker,110,64,1620,23,MINT,True)
    box(80,1007,1760,2,'253148'); txt('SOLANA MOBILE  /  HACKATHON PRODUCT INTRO',80,1030,1550,19,MUTED)
    txt(f'{n:02d} / 06',1720,1028,160,22,MINT,True)

def finish(n,title,time,vo,note):
    assert vo == NARRATION[n-1], f'Narration mismatch on slide {n}'
    s.notes_slide.notes_text_frame.text=f'{time}\n\nENGLISH VOICEOVER\n{vo}\n\nPRODUCTION NOTES\n{note}'
    path=OUT/'slides'/f'{n:02d}.png'; canvas.save(path); images.append(canvas.copy())
    manifest.append(dict(slide=n,title=title,time=time,voiceover=vo,production_notes=note,image=f'slides/{n:02d}.png'))

start(1,'MOVE WITH PURPOSE')
txt('Daily movement.\nVisible progress.',80,210,1120,86,WHITE,True)
txt('A fitness experience built\nfor Solana Mobile.',85,448,1030,45,MUTED)
box(85,628,1000,106); txt('MOVE  /  LEVEL UP  /  DISCOVER',115,662,930,34,MINT,True)
txt('A reason to return, every day.',85,793,1050,37)
box(1318,138,450,800); pic('dashboard.png',1340,155,406,742)
txt('EARLY BUILD · SEP 14, 2026',1340,914,455,21,GOLD,True)
finish(1,'Move with purpose','0:00–0:18',
'What keeps us moving after the first few days? NeonShift turns everyday movement into visible progress, wildlife-inspired shoes, and a reason to return. Built for Solana Mobile, it connects personal achievement with a wider purpose.',
'目前為 2026-09-14 早期實機截圖，含尚未配置鏈上程式的狀態；不是交易證據。正式影片以提交 APK 首頁錄影替換。前 2 秒可加戶外起步，包含在 18 秒內。')

start(2,'SEE YOUR PROGRESS')
txt('Your movement. Your history.',80,163,1760,72,WHITE,True)
txt('NEXT DESIGN · ACTIVITY + OPTIONAL AUTO-SYNC',85,276,1730,27,GOLD,True)
box(80,360,850,510)
txt('ACTIVITY',112,390,780,34,MINT,True)
txt('Month view · Filters · Workout details',112,454,780,31)
for i,line in enumerate(['SEP 17   RUN     5.00 km   30:00','SEP 18   WALK    2.40 km   28:00','SEP 19   RUN     3.00 km   19:00']):
    txt(line,112,548+i*83,780,30,MUTED)
box(990,360,850,510)
txt('AUTO-SYNC WHEN ONLINE',1022,390,780,32,MINT,True)
txt('Your choice: OFF / ON',1022,454,780,34)
txt('Oldest workout first',1022,552,780,39,WHITE,True)
txt('SEP 17  >  SEP 18  >  SEP 19',1022,625,780,32,MINT)
txt('Retry safely. Keep records in order.',1022,724,780,30,MUTED)
txt('DESIGN PREVIEW · SAMPLE DATA · DEVICE VALIDATION PENDING',85,918,1740,24,GOLD,True)
finish(2,'See your progress','0:18–0:43',
'Start with a walk or a run. The app now includes a personal Activity journal with monthly views and workout details. Choose automatic sync when online, with saved workouts processed oldest first. Starter shoes and level upgrades remain free.',
'設計預覽與示例資料，非實機。Activity 月份／篩選／詳情；連網自動同步預設關閉，開啟後依運動開始時間由舊到新。補拍 17→18→19 與較早失敗阻擋後續；不得把摘要同步當成 Health Connect 打卡或自動錢包簽署。')

start(3,'VERIFY A MILESTONE')
txt('Your milestone.\nAn onchain record.',80,165,1710,76,WHITE,True)
txt('Raw health records stay on your phone.',85,382,1730,39,MINT)
steps=[('01','HEALTH SUMMARY','Android Health\nConnect'),('02','VALIDATION','Qualifying activity\nis checked'),('03','WALLET APPROVAL','You approve\nthe claim'),('04','ONCHAIN CLAIM','View the transaction\non Solana devnet')]
for i,(num,head,body) in enumerate(steps):
    x=80+i*450; box(x,505,410,260); txt(num,x+26,528,360,32,MINT,True); txt(head,x+26,594,365,26,WHITE,True); txt(body,x+26,657,365,27,MUTED)
    if i<3: txt('>',x+419,613,30,30,MINT,True)
txt('WORKFLOW PREVIEW · DEVICE VALIDATION PENDING',85,826,1750,26,GOLD,True)
txt('tSKR: devnet test token · No monetary value · Not official SKR',85,896,1750,29,MUTED)
finish(3,'Verify a milestone','0:43–1:29',
'NeonShift connects qualifying daily activity with a claim on Solana devnet. An activity summary from Android Health Connect is checked before you approve the claim in your wallet. The resulting transaction can be viewed in the explorer. Raw health records stay on your phone; they are not published onchain. The prototype uses test tokens with no monetary value. The blockchain records the claim, not every step of your run.',
'本頁為流程圖，非已成功交易。正式影片以合格摘要 → 打卡 → 錢包確認 → 成功 → 同筆 Explorer 實錄覆蓋中間流程區。僅證據完成後改為 REAL DEVICE · DEVNET；縮短等待需標 Wait time shortened。')

start(4,'DISCOVER WILD GUARDIANS')
txt('Progress with a story.',80,163,1760,77,WHITE,True)
txt('NEXT DESIGN · SWITCH EARNED SHOES + MATCHING BACKGROUNDS',85,278,1730,27,GOLD,True)
for x,name,label,sub in [(80,'asian-elephant.png','ASIAN ELEPHANT','Forest theme · Earned shoe selection'),(985,'hawksbill-turtle.png','HAWKSBILL TURTLE','Ocean theme · Background ON / OFF')]:
    box(x,376,855,484); pic(name,x+20,387,815,335); txt(label,x+32,744,790,34,MINT,True); txt(sub,x+32,801,800,28)
txt('DESIGN PREVIEW · COSMETIC ONLY · ACTIVE LEVEL CONTROLS ELIGIBILITY',85,904,1740,24,GOLD,True)
finish(4,'Discover Wild Guardians','1:29–2:03',
'As you level up, discover Wild Guardians, inspired by threatened wildlife. The app now lets you switch between shoes you have earned, with matching forest or ocean backgrounds. Prefer a simpler view? Turn the background off. These choices are cosmetic; your active level still controls eligibility. Each shoe connects your progress with conservation learning.',
'本頁使用專案既有鞋款素材；多鞋切換、森林／海洋背景與關閉開關已有程式；本頁仍是示意，提交版本實機待驗收。正式影片補切兩雙已取得鞋→關背景→故事卡；圖不可冒充實機。使用試拆時標 DEMO PREVIEW，切帳號標 Prepared demo account。不稱隨機 NFT、付費抽獎或保育收益權。')

start(5,'CONNECT WITH COMMUNITY')
txt('From daily movement\nto shared experiences.',80,158,1740,74,WHITE,True)
txt('EVENT PROTOTYPE · WORKFLOW PREVIEW',85,382,1720,27,GOLD,True)
box(80,475,850,280); box(990,475,850,280)
txt('PARTICIPANT',112,508,780,31,MINT,True); txt('Discover an event\nRegister and show a code',112,586,780,42)
txt('AUTHORIZED STAFF',1022,508,780,31,PURPLE,True); txt('Enter the participant code\nConfirm attendance',1022,586,780,42)
txt('NEXT: local community and conservation pilots.',85,825,1740,37)
txt('Device validation pending. Partnerships are planned.',85,902,1740,28,MUTED)
finish(5,'Connect with community','2:03–2:29',
'The journey also reaches beyond solo exercise. This event prototype previews registration and staff-confirmed check-in. Participants show a code; authorized staff confirm attendance. Event participation connects back to your personal journey. Our next step is to test this experience with local communities and conservation partners.',
'雙角色流程示意。完成提交 APK 端到端驗收後替換為報名 → 出示代碼 → 授權 staff 輸入 → 狀態確認，保留 TEST DATA 標示。不可改成 staff 相機掃碼或暗示 NFC 為出席證明；活動報到不等同每日鏈上打卡。')

start(6,'BUILD WHAT COMES NEXT')
txt('Move for yourself.',80,192,1770,88,WHITE,True)
txt('Move together\nfor nature.',80,314,1770,100,MINT,True)
box(80,642,850,200); box(990,642,850,200)
txt('THE EXPERIENCE',112,670,770,27,MINT,True); txt('Movement · Progress\nConservation learning',112,727,770,35)
txt('NEXT: LOCAL PILOTS',1022,670,775,27,PURPLE,True); txt('Planned: paid event services\nSeparate conservation reporting',1022,727,775,33)
txt('NEONSHIFT',85,907,1700,38,WHITE,True)
finish(6,'Build what comes next','2:29–2:50',
'Our direction is simple: move for yourself, move together for nature. We plan to fund useful event services through paying partners, with conservation contributions tracked separately. NeonShift brings movement, meaningful progress, and conservation learning into one mobile experience.',
'最後至少靜止 2 秒。正式發布 APK、原始碼與評審指南後，才在影片說明欄與片尾加可點連結／QR。現有投影片不放未發布連結或佔位網址；尚無已確認合作、捐款或減碳實績。')

prs.save(OUT/'NeonShift_Hackathon_EN.pptx')
images[0].save(OUT/'NeonShift_Hackathon_EN.pdf',save_all=True,append_images=images[1:],resolution=144)
contact=Image.new('RGB',(1440,1215),'#'+BG)
for i,im in enumerate(images): contact.paste(im.resize((720,405)),((i%2)*720,(i//2)*405))
contact.save(OUT/'overview.jpg',quality=93)
(OUT/'slide-content.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
(OUT/'voiceover-v3-archive.txt').write_text('\n\n'.join(v['voiceover'] for v in manifest)+'\n')
(OUT/'preview.html').write_text('<!doctype html><html lang="en"><meta charset="utf-8"><title>NeonShift — Hackathon</title><style>body{margin:0;background:#070b16;color:#a8b6cf;font:16px Arial}main{max-width:1280px;margin:auto}img{display:block;width:100%;margin:20px 0 8px}p{margin:0 0 32px}@media print{img{break-before:page;margin:0}p{display:none}@page{size:landscape;margin:0}}</style><main>'+''.join(f'<img src="{v["image"]}" alt="{escape(v["title"])}"><p>{v["slide"]:02d} / {v["time"]} — {escape(v["title"])}</p>' for v in manifest)+'</main></html>')
print('Built 6 slides, PPTX, PDF, PNGs, preview, overview, narration, and manifest.')

# Rebuild historical v3 visuals only; current v4 narration is maintained separately.
sections=['# NeonShift｜逐頁英文文案與中文分鏡\n\n> 歷史 v3 六段視覺參考；現行五段配音／時間軸見 ../../docs/store/demo-video.md。勿與 voiceover-en.txt 混用。\n\n2026-09-22 更新；第 2、4 頁仍為示意素材，功能已有程式，提交版實機驗收待完成。']
for slide_obj, item in zip(prs.slides, manifest):
    visible='\n\n'.join('\n'.join('> '+line for line in sh.text.splitlines()) for sh in slide_obj.shapes if sh.has_text_frame and sh.text.strip())
    sections.append(f"## {item['slide']:02d}｜{item['title']}｜{item['time']}\n\n### 投影片畫面文字\n\n{visible}\n\n### English voiceover\n\n{item['voiceover']}\n\n### 中文分鏡與製作註記\n\n{item['production_notes']}")
(OUT/'storyboard.md').write_text('\n\n'.join(sections)+'\n')
