"""Build the eleven-scene FlexClip v4 import pack from canonical narration."""
from pathlib import Path
from html import escape
import json, csv, re, argparse, shutil
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
BASE=Path(__file__).resolve().parent
ROOT=BASE.parents[2]
parser=argparse.ArgumentParser()
parser.add_argument('--lang', choices=['en','zh-TW'], default='en')
LANG=parser.parse_args().lang
OUT=BASE if LANG=='en' else BASE/'zh-TW'
OUT.mkdir(exist_ok=True)
LOCAL=json.loads((BASE/'zh-TW-content.json').read_text()) if LANG=='zh-TW' else None
BG='070B16'; PANEL='111A2B'; WHITE='F3F7FF'; MUTED='A8B6CF'; MINT='35E5C5'; PURPLE='AD8CFF'; GOLD='FFD080'
FONT_DIR=Path('/System/Library/Fonts/Supplemental')
prs=Presentation(); prs.slide_width=Inches(13.333333); prs.slide_height=Inches(7.5)
(OUT/'slides').mkdir(exist_ok=True)
(OUT/'narration').mkdir(exist_ok=True)
paras=(ROOT/'docs/store/demo-voiceover-en.txt').read_text().strip().split('\n\n')
assert len(paras)==5
# Split each chapter at sentence boundaries, without rewriting the approved script.
# 每章的切點（句數）；E 章拆三頁：保育學習／SKR 質押加成與保育提撥（規劃）／片尾
cuts=[[2],[2],[3],[3],[3,6]]
voices=[]
for p,ks in zip(paras,cuts):
    sentences=re.split(r'(?<=[.!?]) +', p)
    bounds=[0]+ks+[len(sentences)]
    voices += [' '.join(sentences[a:b]) for a,b in zip(bounds,bounds[1:])]
assert ' '.join(voices)==' '.join(paras)
def font(size,bold=False):
    if LANG=='zh-TW':
        return ImageFont.truetype('/System/Library/Fonts/STHeiti Medium.ttc' if bold else '/System/Library/Fonts/STHeiti Light.ttc',size)
    path=FONT_DIR/('Arial Bold.ttf' if bold else 'Arial.ttf')
    if not path.exists():
        path=Path('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf' if bold else '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')
    return ImageFont.truetype(str(path),size)

def box(x,y,w,h,color=PANEL):
    shape=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x/144), Inches(y/144), Inches(w/144), Inches(h/144))
    shape.fill.solid(); shape.fill.fore_color.rgb=RGBColor.from_string(color); shape.line.fill.background()
    draw.rectangle((x,y,x+w,y+h),fill='#'+color)

def txt(t,x,y,w,size=36,color=WHITE,bold=False):
    lines=t.split('\n'); lh=round(size*1.23); h=lh*len(lines)+8
    assert x+w<=1920 and y+h<=1080, (t,x,y,w,h)
    tb=s.shapes.add_textbox(Inches(x/144),Inches(y/144),Inches(w/144), Inches(h/144))
    tf=tb.text_frame; tf.word_wrap=False; tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(lines):
        assert draw.textlength(line,font=font(size,bold)) <= w, ('text too wide',line,w)
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.space_after=Pt(0); p.space_before=Pt(0); p.line_spacing=Pt(lh/2)
        r=p.add_run(); r.text=line; r.font.name='Heiti TC' if LANG=='zh-TW' else 'Arial'; r.font.size=Pt(size/2); r.font.bold=bold; r.font.color.rgb=RGBColor.from_string(color)
        draw.text((x,y+i*lh),line,font=font(size,bold),fill='#'+color,anchor='lt')

def pic(name,x,y,w,h):
    path=ROOT/name; im=Image.open(path).convert('RGB'); ratio=min(w/im.width,h/im.height)
    iw,ih=round(im.width*ratio),round(im.height*ratio); xx=x+(w-iw)//2; yy=y+(h-ih)//2
    canvas.paste(im.resize((iw,ih),Image.Resampling.LANCZOS),(xx,yy))
    s.shapes.add_picture(str(path),Inches(xx/144),Inches(yy/144),width=Inches(iw/144),height=Inches(ih/144))


scenes=[
 ('A','Daily movement.\nVisible progress.','A reason to return, every day.','MOVE  /  GROW  /  COLLECT',12,'output/hackathon-flexclip-en/captures/2026-10-02/2026-10-02_c14_s01_home-top.png','REAL DATA · NO DEMO OVERRIDE','10/2 v14 實機首頁（真實資料，狀態列以 demo mode 固定）。最終 APK 外觀相同即可沿用。'),
 ('A','Built for\nSolana Mobile.','Phone-based workouts.\nHealth summaries. Wallet approval.','MOBILE-FIRST FITNESS',13,None,'PRODUCT OVERVIEW','具體說明手機運動、健康摘要與錢包用途；右側為功能圖解，不是實機。'),
 ('B','Walk. Run.\nSee your progress.','Record → Save → Review','YOUR ACTIVITY JOURNAL',17,'output/hackathon-flexclip-en/captures/2026-10-02/2026-10-02_c14_s04_summary-finished.png','REAL DATA · NO DEMO OVERRIDE','10/2 v14 實機：9/30 跑步 3.89 km 的運動摘要（已保存、已同步）。'),
 ('B','Sync when\nyou are ready.','Optional online sync.\nOldest workout first.','STARTER SHOES + UPGRADES: FREE',18,'output/hackathon-flexclip-en/captures/2026-10-02/2026-10-02_c14_s04_profile-sync.png','REAL DATA · NO DEMO OVERRIDE','10/2 v14 實機：Auto-sync 開、0 pending、最後成功 9/30——真實同步狀態。保留「保存不等於獎勵資格」的旁白。'),
 ('C','Your approval.\nAn onchain claim.','Qualifying summary\n→ Wallet approval → Explorer','SOLANA DEVNET',18,None,'WORKFLOW OVERVIEW','黑客松 Demo 在本頁嵌入合格摘要、錢包與同筆 Explorer 實錄。'),
 ('C','Collect milestones.\nKeep the meaning.','Achievements have their own\nrequirements and approval.','RAW HEALTH RECORDS: NOT ONCHAIN',17,'output/hackathon-flexclip-en/captures/2026-10-02/2026-10-02_c14_s06_nft-first5k-detail.png','REAL DATA · NO DEMO OVERRIDE','10/2 v14 實機：First 5K 成就收藏詳情（9/22 鑄造、devnet）。這是成就收藏，不是付款結果。'),
 ('D','Earn eligibility.\nUnlock your frame.','Verified first 5K eligibility.\nServer-set order, price, recipient.','COSMETIC PAYMENT',22,'docs/evidence/2026-10-01-skr-devnet-eligible.png','TEST SKR · DEVNET · NOT OFFICIAL SKR','10/1 v11 實機：資格通過後的訂單畫面（2.5 SKR 固定價、伺服器指定收款人）。與第 8 頁同一帳號、同一包。'),
 ('D','Confirmed payment.\nUnlocked frame.','Wallet approval → Onchain check\n→ Cosmetic unlock','NO XP, RECORDS OR ACHIEVEMENTS BOUGHT',23,'docs/evidence/2026-10-01-skr-devnet-paid.png','TEST SKR · DEVNET · NOT OFFICIAL SKR','10/1 v11 實機付款成功截圖（簽章見 docs/evidence/2026-10-01-skr-devnet-payment.md）。10/2 決定沿用此圖，不另錄付款影片。'),
 ('E','Progress with\na wildlife story.','Conservation learning today.\nCommunity pilots are planned.','NEXT: LOCAL PILOTS',12,'output/hackathon-flexclip-en/captures/2026-10-02/2026-10-02_c14_s09_elephant-habitat.png','IN-APP STORY CARD · PILOTS PLANNED','10/2 v14 實機：亞洲象物種故事卡。物種資料引用 WWF，不代表合作或捐款。'),
 ('E','Next: stake SKR.\nBoost by level.','Higher shoe levels earn higher boosts.\nA fixed share of staking rewards\nfunds wildlife conservation.','ROADMAP  ·  NOT LIVE  ·  RATES NOT FINAL',13,None,'PLANNED · NOT IMPLEMENTED','規劃中：質押 SKR 提升任務獎勵，等級越高加成越高；質押獎勵固定比例撥入獨立追蹤的保育基金。右側比率為示意，未定案；不承諾收益，主網與官方 SKR 尚未啟用。'),
 ('E','Move for yourself.\nMove together\nfor nature.','Explore NeonShift.','ANDROID APP  /  SOURCE  /  REVIEWER GUIDE',12,None,'NEONSHIFT · SOLANA MOBILE','最後至少靜止 2 秒。公開連結放影片說明欄，不使用佔位 URL。'),
]
images=[]; manifest=[]; elapsed=0
for i,(chapter,title,body,tag,duration,photo,label,note) in enumerate(scenes,1):
    if LOCAL:
        item=LOCAL[i-1]
        title,body,tag,label=(item[k] for k in ['title','body','tag','label'])
    s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=RGBColor.from_string(BG)
    canvas=Image.new('RGB',(1920,1080),'#'+BG); draw=ImageDraw.Draw(canvas)
    box(80,65,8,25,MINT); txt('NEONSHIFT  /  '+chapter+('  ·  中文對稿' if LOCAL else ''),110,65,1100,26,MINT,True)
    txt(f'{i:02d} / {len(scenes)}',1700,65,150,26,MUTED)
    width=1120 if photo else 1760
    txt(title,80,190,width,70,WHITE,True)
    txt(body,85,470,width,36,MUTED)
    txt(tag,85,650,width,25,MINT,True)
    if photo:
        box(1280,152,540,650)
        pic(photo,1300,170,500,610)
    diagrams={
        2:['PHONE WORKOUTS','HEALTH SUMMARY','WALLET APPROVAL'],
        5:['QUALIFYING SUMMARY','WALLET APPROVAL','DEVNET EXPLORER'],
    }
    diagrams_zh={2:['手機運動紀錄','健康資料摘要','錢包核准'],5:['符合條件的摘要','使用錢包核准','Devnet 鏈上查證']}
    # 第 10 頁：各等級質押加成（示意）＋保育提撥
    if i==10:
        rows=[('LV.1  ORIGIN','+5%'),('LV.2','+10%'),('LV.3','+15%'),('LV.4','+20%'),('LV.5','+25%')]
        rows_zh=[('LV.1 原點','+5%'),('LV.2','+10%'),('LV.3','+15%'),('LV.4','+20%'),('LV.5','+25%')]
        txt('質押加成（示意）' if LOCAL else 'STAKING BOOST (ILLUSTRATIVE)',1230,160,600,24,MUTED,True)
        for j,(lv,rate) in enumerate(rows_zh if LOCAL else rows):
            y=205+j*92
            box(1230,y,600,76)
            txt(lv,1260,y+22,330,27,WHITE,True)
            txt(rate,1700,y+22,110,27,MINT,True)
        y=205+5*92+14
        box(1230,y,600,90,'1B2A3F')
        txt('保育提撥' if LOCAL else 'CONSERVATION SHARE',1260,y+14,420,24,GOLD,True)
        txt('質押獎勵的 10%' if LOCAL else '10% of staking rewards',1260,y+50,520,24,WHITE)
    if i in diagrams:
        for j,line in enumerate((diagrams_zh if LOCAL else diagrams)[i]):
            y=235+j*172
            box(1230,y,600,130)
            txt(f'{j+1:02d}',1260,y+20,90,27,MINT,True)
            txt(line,1260,y+65,535,25,WHITE,True)
    # 實機截圖一律標出拍攝日與 APK 版本代碼，讓評審分得出來源
    caps={1:('OCT 02','10/2','14'),3:('OCT 02','10/2','14'),4:('OCT 02','10/2','14'),6:('OCT 02','10/2','14'),7:('OCT 01','10/1','11'),8:('OCT 01','10/1','11'),9:('OCT 02','10/2','14')}
    if i in caps:
        en_d,zh_d,code=caps[i]
        txt(f'實機截圖 · {zh_d} · APK 版本代碼 {code}' if LOCAL else f'DEVICE CAPTURE · {en_d} · APK CODE {code}',85,746,1120,25,GOLD)
    txt(label,85,843,1750,25,GOLD,True)
    # Bottom 150 pixels are deliberately clear for FlexClip subtitles.
    voice=LOCAL[i-1]['voiceover'] if LOCAL else voices[i-1]
    s.notes_slide.notes_text_frame.text=voice
    file=f'slides/{i:02d}.png'; canvas.save(OUT/file); images.append(canvas.copy())
    (OUT/'narration'/f'{i:02d}.txt').write_text(voice+'\n')
    manifest.append(dict(scene=i,chapter=chapter,start_seconds=elapsed,duration_seconds=duration,title=title.replace('\n',' '),voiceover=voice,image=file,source=photo,status=label,production_notes=note))
    elapsed+=duration
assert elapsed==177, elapsed
prs.save(OUT/'NeonShift_FlexClip_v4.pptx')
suffix='EN' if LANG=='en' else 'ZH-TW_Review'
shutil.copyfile(OUT/'NeonShift_FlexClip_v4.pptx',BASE/f'NeonShift_FlexClip_v4_{suffix}.pptx')
images[0].save(OUT/'NeonShift_FlexClip_v4.pdf',save_all=True,append_images=images[1:],resolution=144)
shutil.copyfile(OUT/'NeonShift_FlexClip_v4.pdf',BASE/f'NeonShift_FlexClip_v4_{suffix}.pdf')
contact=Image.new('RGB',(1152,324*((len(images)+1)//2)),'#'+BG)
for i,im in enumerate(images): contact.paste(im.resize((576,324)),((i%2)*576,(i//2)*324))
contact.save(OUT/'overview.jpg',quality=90)
(OUT/'scene-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
with (OUT/'timeline.csv').open('w',newline='') as f:
    w=csv.writer(f); w.writerow(['scene','chapter','start_seconds','duration_seconds','voiceover_file'])
    for m in manifest: w.writerow([m['scene'],m['chapter'],m['start_seconds'],m['duration_seconds'],f"narration/{m['scene']:02d}.txt"])
(OUT/'narration-all.txt').write_text('\n\n'.join(m['voiceover'] for m in manifest)+'\n')
(OUT/'scene-guide.md').write_text('# FlexClip 逐頁貼稿與時間表\n\n英文版旁白用於 FlexClip；中文版僅供對稿，不代表中文朗讀時長。時間為英文剪輯目標。\n\n'+'\n\n'.join(f"## {m['scene']:02d} · {m['title']}\n\n{m['start_seconds']}–{m['start_seconds']+m['duration_seconds']} 秒 · {m['status']}\n\n### Voiceover / 旁白\n\n{m['voiceover']}\n\n### 製作備註（不要配音）\n\n{m['production_notes']}" for m in manifest)+'\n')
(OUT/'preview.html').write_text('<!doctype html><html lang="en"><meta charset="utf-8"><title>NeonShift FlexClip v4</title><style>body{background:#070b16;color:#f3f7ff;font:18px Arial;max-width:1100px;margin:30px auto}img{width:100%}section{margin-bottom:40px}</style><h1>NeonShift · FlexClip v4 · 177 second target</h1>'+''.join(f'<section><img src="{m["image"]}" alt="{escape(m["title"])}"><p>{m["scene"]:02d} · {m["duration_seconds"]} s · {escape(m["voiceover"])}</p></section>' for m in manifest)+'</html>')
print(f'Built {LANG}: {len(scenes)} slides; English target {elapsed} seconds; aligned chapter mapping.')

if LOCAL:
    en=json.loads((BASE/'scene-manifest.json').read_text())
    review=['# NeonShift 中英逐頁對稿\n\n英文版用於 FlexClip；中文版僅供審稿。兩版頁碼、素材與英文時間軸相同；中文不作配音時長承諾。App 原圖維持原始介面語言。']
    for e,z in zip(en,manifest):
        review.append(f"## {e['scene']:02d} · {e['title']} / {z['title']}\n\n英文時間：{e['start_seconds']}–{e['start_seconds']+e['duration_seconds']} 秒\n\n**English**\n\n{e['voiceover']}\n\n**中文意涵**\n\n{z['voiceover']}\n\n**素材界線**：{z['status']}。{z['production_notes']}")
    (BASE/'bilingual-review.md').write_text('\n\n'.join(review)+'\n')
