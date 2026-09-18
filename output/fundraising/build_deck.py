# -*- coding: utf-8 -*-
"""Build editable PPTX and matching HTML review copy. No remote runtime calls."""
from pathlib import Path
from html import escape
import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
BG='050711'; PANEL='0B1020'; WHITE='F4F8FF'; MUTED='AAB7CC'; MINT='30EBC8'; VIOLET='9B6CFF'; CYAN='24C8FF'; AMBER='FFCB66'
FONT='Microsoft JhengHei'
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
html_slides=[]; outline=[]; current=[]
SOURCES={
 'S1':('Global Wellness Institute｜2025 Global Wellness Economy Monitor','https://globalwellnessinstitute.org/wp-content/uploads/2025/11/2025-GWI-WE-Monitor_DIGITAL-FINAL.pdf'),
 'S2':('HFA｜Fitness Industry’s Global Momentum Continued in 2024','https://www.healthandfitness.org/fitness-industrys-global-momentum-continued-in-2024-new-report-shows/'),
 'S3':('Solana Mobile｜SKR is Live，2026-07-14','https://solanamobile.com/blog/skr-is-live')}

def rect(s,x,y,w,h,color=PANEL,line=None):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
 sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor.from_string(color)
 sh.line.fill.background()
 if line: sh.line.color.rgb=RGBColor.from_string(line)
 current.append('<div class="box" style="left:%sin;top:%sin;width:%sin;height:%sin;background:#%s;%s"></div>'%(x,y,w,h,color,('border:1px solid #'+line) if line else ''))
 return sh

def text(s,t,x,y,w,h,size=22,color=WHITE,bold=False,url=None):
 tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
 tf=tb.text_frame; tf.word_wrap=True
 tf.margin_left=tf.margin_right=Inches(.01); tf.margin_top=tf.margin_bottom=Inches(.01)
 for i,line in enumerate(t.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.space_after=Pt(8)
  run=p.add_run(); run.text=line; run.font.name=FONT; run.font.size=Pt(size); run.font.bold=bold; run.font.color.rgb=RGBColor.from_string(color)
  ea=OxmlElement('a:ea'); ea.set('typeface',FONT); run._r.get_or_add_rPr().append(ea)
  if url: run.hyperlink.address=url
 current.append('<div class="text" style="left:%sin;top:%sin;width:%sin;height:%sin;font-size:%spt;color:#%s;font-weight:%s">%s</div>'%(x,y,w,h,size,color,700 if bold else 400,escape(t).replace('\n','<br>')))
 return tb

def pic(s,name,x,y,w,h):
 path=ROOT/'assets/brand'/name
 s.shapes.add_picture(str(path),Inches(x),Inches(y),width=Inches(w),height=Inches(h))
 current.append('<img style="left:%sin;top:%sin;width:%sin;height:%sin" src="../../assets/brand/%s">'%(x,y,w,h,name))

def slide(title,kicker='投資人簡報草稿',sub=None,source=None):
 global current
 current=[]
 s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=RGBColor.from_string(BG)
 rect(s,.48,.39,.08,.23,MINT)
 text(s,'NEONSHIFT  /  '+kicker,.72,.32,11,.35,11,MINT,True)
 text(s,title,.72,.93,11.9,.78,30,WHITE,True)
 if sub: text(s,sub,.74,1.83,11.7,.7,16,MUTED)
 text(s,'募資討論稿・2026.09.18｜永續策略／商業與數值假設待驗證',.74,7.07,10,.22,9,MUTED)
 text(s,'%02d'%len(prs.slides),12.02,7.03,.5,.3,12,MINT)
 if source:
  title_,url=SOURCES[source]; text(s,source+'｜'+title_,.74,6.7,11.7,.25,9,MUTED,url=url)
 return s

def finish(s,title,notes):
 s.notes_slide.notes_text_frame.text=notes
 html_slides.append('<section>'+''.join(current)+'</section>')
 outline.append({'slide':len(prs.slides),'title':title,'notes':notes})

def cards(s,items,y=2.75,h=3.35):
 n=len(items); gap=.22; w=(11.85-gap*(n-1))/n
 for i,(head,body) in enumerate(items):
  x=.74+i*(w+gap); rect(s,x,y,w,h)
  text(s,'%02d'%(i+1),x+.23,y+.22,w-.46,.35,13,MINT)
  text(s,head,x+.23,y+.79,w-.46,.58,22,WHITE,True)
  text(s,body,x+.23,y+1.55,w-.46,h-1.7,17,MUTED)

s=slide('為自己而動，為棲地同行',kicker='運動 × 保育 × 可持續合作')
text(s,'NeonShift 賽博躍遷',.75,2.1,7,.7,34,MINT,True)
text(s,'把日常運動、在地保育與合作活動，\n串成能持續參與、可查核的行動。',.75,3.1,7,1.1,24)
text(s,'從 Solana Mobile 社群切入\n面向跑團、主辦方與運動品牌的合作平台',.75,4.6,6.6,1,18,MUTED)
pic(s,'neonshift-splash-v1.png',9.5,1.8,2.15,4.66)
finish(s,'投資主張','2026-09-18 策略提案：日常運動 × 在地保育參與 × 可查核合作成果。共同主張「為自己而動，為棲地同行」。從 Android／Seeker 原型切入，優先試辦品牌付費的四週行動季；SOL 指 Solana／SOL 生態。畫面為設計稿，保育合作、付費客戶與成果尚待驗證。完整研究：docs/sustainability-direction.md。')

s=slide('活動結束後，關係不應歸零',sub='待訪談驗證的三個痛點：參與者難持續、主辦方難回訪、品牌難衡量。')
cards(s,[('參與者','健康紀錄與活動成績分散\n缺乏回到下一次活動的理由'),('活動主辦方','宣傳、報到、發放與成績分開\n難追蹤跨場次參與'),('合作品牌','一次曝光之後缺少後續連結\n需要看見到場與實際領取')])
finish(s,'問題假設','本頁為團隊待驗證的客戶問題假設，非市場調查結論。建議訪談至少 10 位主辦方／品牌及 20 位參與者，先判斷是否願意付費解決。')

s=slide('建立「日常任務 → 活動 → 再參與」循環',sub='核心價值是提升參與與回訪；獎勵是體驗的一部分。')
cards(s,[('日常養成','步數與睡眠任務\n成長型跑鞋與進度回饋'),('線下參與','合作活動宣傳與報名\nNFC／QR 報到、贈品與徽章'),('持續連結','活動成績冊與個人歷程\n下一場邀請與回訪衡量')])
finish(s,'產品價值循環','既有 BRD 包含日常健康任務與新增合作活動。再參與與品牌邀請屬商業體驗提案，需依同意及通知政策實作。無須向投資人描述後端或鏈上指令。')

s=slide('用同一個品牌體驗，承接線上與現場',sub='深藍黑底、青綠至紫色的 N 躍遷識別；可延伸至活動、NFC 載具與成績憑證。')
pic(s,'neonshift-identity-v1.png',.75,2.65,6.45,4.3)
pic(s,'neonshift-loading-v1.png',8.5,2.05,2.05,4.44)
text(s,'既有視覺設計稿\n非正式上線截圖',10.85,4.3,1.55,1,13,MUTED)
finish(s,'品牌與產品展示','使用 assets/brand 既有設計稿。Logo、Splash、Loading 已完成概念設計；此頁不得被解讀為所有產品功能已完成或取得品牌合作。')

s=slide('合作活動，是第一個可收費的切入口',sub='案例示意：跑團／品牌共同辦理一場 300 人活動，並追蹤後續回訪。')
cards(s,[('活動前','共同宣傳與來源追蹤\n報名名單與參加規則'),('活動中','NFC／QR 到場核驗\n庫存控管、贈品與數位徽章'),('活動後','主辦方確認成績\n成果報告與下一場參與')])
finish(s,'合作活動模式','300 人為流程示意，不是已舉辦活動或已取得客戶。第一階段發放实體贈品及鏈下數位徽章；NFC 標籤只作入口，不保證不可複製。活動成績由來源方確認，不冒充健康資料證明。')

s=slide('產業規模足夠大，但必須聚焦付費場景',kicker='產業規模｜2024 年實績',sub='全球身體活動經濟包含運動服務、用品與科技；數字不能當成平台可取得的營收。',source='S1')
for i,(num,label,desc) in enumerate([('1.14 兆','美元｜全球身體活動','2024 年：1,143.9 十億美元'),('859 億','美元｜全球健身科技','包含硬體、軟體及服務'),('3,197 億','美元｜亞太身體活動','區域背景，非活動軟體市場')]):
 x=.74+i*4.03; rect(s,x,2.9,3.8,3.1); text(s,num,x+.2,3.24,3.4,.8,39,MINT,True);text(s,label,x+.2,4.28,3.4,.6,18);text(s,desc,x+.2,5.12,3.4,.65,13,MUTED)
finish(s,'產業規模','來源 S1，2025 年報告，2024 年實績；身體活動章節，PDF p.61–65。健身科技 85.9 十億美元、亞太 319.7 十億美元、全球 1,143.9 十億美元。三者有包含／重疊关系，不相加。健身科技包含硬體，不是 NeonShift 的 TAM。查核日期 2026-09-14。')

s=slide('數位運動正在成長，現場經營仍需驗證',kicker='市場動能',sub='外部資料支持需求背景；NeonShift 的商業可行性仍要用付費試辦證明。',source='S1')
cards(s,[('12.1%','2024 年全球健身科技\n較前一年成長率 [S1]'),('21.5%','健身 App／串流／隨選服務\n2019–2024 年複合成長率 [S1]'),('社群切入','先以 Seeker 社群測試參與\n再驗證主辦方與品牌付費')])
finish(s,'市場動能','S1 PDF p.63–64：科技市場同比成長 12.1%；App、串流與隨選服務支出五年 CAGR 21.5%。不代表本公司預期增速。補充 S2：有可比資料國家在 2024 年營收平均增 8%，非全球加權成長率。S3 官方 2026-07-14 公告逾 10 萬人具 SKR 領取資格，並非活躍數、健身用戶數或本公司觸達人數。')

s=slide('從可接觸的主辦方，推算收入機會',kicker='市場估算｜情境而非市場普查',sub='建議先驗證台灣跑團／活動合作，再評估跨區擴張；地區選擇為待確認策略。')
cards(s,[('產業上界','全球健身科技：859 億美元\n含硬體，僅作參考 [S1]'),('服務市場情境','假設 1,000 家合格主辦方\n每年 6 場 × 每場 NT$30,000\n＝ NT$1.8 億／年'),('早期取得情境','20 家付費主辦方\n每年 4 場 × 每場 NT$30,000\n＝ NT$240 萬／年')])
finish(s,'可服務市場與初期機會','SAM 尚無可核驗的區域客戶母體，因此此頁不宣稱正式 TAM/SAM/SOM。1,000 家是假設，不是已盤點家數；須建立活動頻率、預算與 Android 參與者適配的名單。20 家是第 3 年基準情境，非已簽約或預測。年度場次收入不是 ARR；不同幣別未換算或相加。來源 S1 僅支持產業上界。')

s=slide('先讓主辦方與品牌付費，再擴大會員價值',kicker='商業模式提案',sub='收費依服務價值；不將代幣升值或出售健康資料當作營收來源。')
cards(s,[('活動服務費','每場 NT$30,000 起（假設）\n報名、核銷、成績與報告\n硬體／現場人力另報價'),('品牌合作專案','每案 NT$100,000（假設）\n主題任務與成效彙總\n獎品與媒體費另外編列'),('後續訂閱','多場活動管理與進階報告\n需先驗證反覆使用與續約\n尚未納入收入模型')])
finish(s,'商業模式','為本次募資簡報提出的商業假設，未更動既有產品範圍／價目表。價格未完成客戶訪談。活動與品牌專案須有獨立可計費工作範圍，同一服務不得重複計收入；財務模型採平台服務收入，不含代收贊助款、獎品及售票總額。')

s=slide('讓每一場活動，成為下一場的獲客入口',kicker='市場進入策略｜建議')
cards(s,[('01｜共同試辦','先找 3 家設計合作夥伴\n跑團、主辦方、運動品牌\n確認痛點與付費意願'),('02｜可複製交付','統一報到／發放／成績流程\n衡量每場交付成本\n以真實案例取得續辦'),('03｜跨場次回訪','追蹤同意參與者的再報名\n共同品牌導流\n先擴同客戶，再擴地區')])
finish(s,'市場進入','3 家為下一階段招募目標，尚非已確認合作。初期 Android／Seeker 限制可服務人群；公開網頁入口不等於已具跨平台完整參加體驗。應驗證限制是否妨礙主辦方採用，再決定是否正式變更平台範圍。')

s=slide('差異來自跨場次經營，而非單一 NFC 功能',kicker='定位與替代方案',sub='以下比較的是使用情境，不是宣稱所有競品都缺少某項功能。')
cards(s,[('既有替代方案','運動紀錄 App：日常數據\n活動工具：單場交付\n品牌行銷：短期觸及'),('NeonShift 提案','日常任務＋合作活動＋成績冊\n把報到、領取與回訪串連\n提供可衡量合作成果'),('待建立的優勢','穩定交付與主辦方續約\n經同意的跨場參與歷程\n合作密度與品牌辨識')])
finish(s,'定位','本頁為自家定位分析，不是逐一競品功能研究。NFC、積分與徽章容易被複製，不單獨視為護城河。應把合作續約、交付效率、跨活動回訪視為驗證指標，且不能以鎖住用戶資料建立優勢。')


s=slide('讓新玩家看懂玩法，讓努力值得紀念',kicker='產品體驗｜已實作，待實機驗收',sub='新手指南 → 每日任務 → 裝備成長 → 成就收藏；從理解玩法到感受榮耀。')
cards(s,[('第一次就知道怎麼玩','四張玩法卡＋首次四步\n可略過、個人頁可重看\n繁中／英文同步'),('取得 NFT 的榮耀感','卡背翻轉、金色光環\n粒子與卡面掃光\n新領取成功才揭曉'),('升降等都有回饋','升等：能量擴散與進場\n降等：冷色收縮與下沉\n支援減少動態效果')])
finish(s,'新手與成就體驗','2026-09-15 自動驗證：動畫相關 17 項、新手與既有流程 23 項測試通過，型別檢查通過；是分次驗證，不當作完整端到端驗收。NftReveal 尚為鞋／徽章展示，非所有 metadata 圖片。低階裝置效能、讀屏與理解度待實機。來源：PG-A-17、PG-UX-01、Style §25。')

# Wildlife collection: rendered SVG motifs are design illustrations, not device screenshots.
def wildlife_pic(s,level,variant,x,y,w,h):
 name=f'lv{level}-{variant}.png'
 s.shapes.add_picture(str(OUT/'wildlife-assets'/name),Inches(x),Inches(y),width=Inches(w),height=Inches(h))
 current.append(f'<img style="left:{x}in;top:{y}in;width:{w}in;height:{h}in" src="wildlife-assets/{name}">')

WILDLIFE_SOURCES={
 'W1':('WWF｜亞洲象','https://www.worldwildlife.org/species/elephant/asian-elephant/'),
 'W2':('WWF｜玳瑁','https://www.worldwildlife.org/species/sea-turtle/hawksbill-turtle/'),
 'W3':('WWF｜老虎','https://www.worldwildlife.org/species/tiger/'),
 'W4':('WWF｜遠東豹','https://www.worldwildlife.org/species/amur-leopard/')}

s=slide('荒野守護：每次升階，認識一種生命',kicker='跑鞋設計｜App 已實作，待實機美術驗收',sub='Lv.1 保留普通「原點」；Lv.2 起將動物特徵融入鞋款。遊戲等級不代表物種價值或瀕危程度。')
for i,(name,detail,topic,tint) in enumerate([
 ('森行・亞洲象','象耳護片 × 象鼻曲線','棲地破碎與人象共存',CYAN),
 ('潮盾・玳瑁','重疊龜甲 × 鰭狀後跟','珊瑚礁與海龜保護',VIOLET),
 ('林焰・老虎','漸尖虎斑 × 森林光軌','森林、獵物與反盜獵','FF4FD8'),
 ('雪影・遠東豹','斷環豹紋 × 冬毛鞋頭','森林廊道與棲地連通',MINT)]):
 x=.74+i*3.02
 rect(s,x,2.75,2.8,3.66,line=tint)
 text(s,f'LV.{i+2}  /  WILD GUARDIANS',x+.15,2.98,2.5,.32,10,tint,True)
 text(s,name,x+.15,3.43,2.5,.5,20,WHITE,True)
 wildlife_pic(s,i+2,'dawn',x+.1,4.02,2.6,1.6)
 text(s,detail,x+.15,5.7,2.5,.32,12,tint)
 text(s,topic,x+.15,6.09,2.5,.26,11,MUTED)
text(s,'SVG 設計示意，非實機截圖；數位仿生材質，不使用動物皮毛、象牙或玳瑁。',.76,6.66,11.7,.28,11,MUTED)
finish(s,'荒野守護跑鞋','首版系列 wild-guardians-v1，保留原等級與 XP 規則；首頁、裝備、藝廊共用動物圖層。鞋圖由既有 SVG 與新動物元素產生，是簡報設計示意而非實機畫面。App 已實作，待實機辨識度與旋轉驗收。物種故事來源 W1–W4，2026-09-17 查核；不宣稱保育合作。規格 docs/design/wild-guardian-shoes.md。')

s=slide('同一階跑鞋，也有屬於你的細節',kicker='成長盲盒｜外觀玩法第一版',sub='達到鞋階 → 揭曉細節款 → 保留個人樣式；從每日進度延伸到收藏與分享。')
for i,(variant,name,detail,tint) in enumerate([
 ('dawn','晨曦縫光','金色縫線與動物紋樣',AMBER),
 ('dusk','暮色紋跡','紫色紋樣與虛線軌',VIOLET),
 ('aurora','極光軌線','薄荷紋樣與加粗光軌',MINT)]):
 x=.74+i*4.03
 rect(s,x,2.76,3.8,3.25,line=tint)
 text(s,name,x+.2,2.98,3.4,.45,22,tint,True)
 wildlife_pic(s,4,variant,x+.2,3.57,3.4,2.09)
 text(s,detail,x+.2,5.64,3.4,.28,13,MUTED)
text(s,'同錢包固定款式  ·  僅外觀差異  ·  不增加能力  ·  無付費重抽',.8,6.27,11.6,.4,18,MINT,True)
finish(s,'盲盒微差異','展示同一 Lv.4 老虎的三款細節，並非三個不同能力等級。未解鎖詳情顯示成長盒，升階沿用揭曉流程。現行依系列 ID＋錢包＋等級固定雜湊分配，可預測，非鏈上隨機抽取或稀有度證明；不宣稱精確抽中機率。款式為 App 外觀，與既有鏈上紀念 NFT 分開。重開、重裝與跨裝置同錢包保持相同結果；未建立二級交易或付費重抽。')

s=slide('從收藏一雙鞋，到認識牠的棲地',kicker='保育教育 × 系列延伸',sub='把保育故事放進鞋款介紹與升階時刻；將未來合作建立在可辨識的系列內容上。')
cards(s,[('現在｜保育介紹','中英雙語物種故事\n棲地、威脅與一個小行動\n附可開啟的 WWF 來源'),('下一步｜系列共創','品牌／創作者聯名提案\n主題鞋面與活動收藏\n授權、合作與切換待完成'),('未來｜資產與成效','獨立款式 NFT metadata\n可追溯的款式取得紀錄\n回訪與閱讀成效待驗證')],y=2.65,h=3.38)
for i,(key,(label,url)) in enumerate(WILDLIFE_SOURCES.items()):
 text(s,label,.8+i*3.02,6.23,2.8,.3,11,MINT,url=url)
text(s,'保育資料查核：2026.09.17；目前沒有已成立的保育合作、捐款或聯名。',.8,6.68,11.6,.27,11,MUTED)
finish(s,'保育教育與系列路線','已完成 App 故事卡、來源連結及系列識別；聯名發行、系列切換、NFT 獨立款式 metadata 尚未實作。現在的共用 kind URI 無法表示每個錢包的獨立款式，需先完成每資產權威紀錄。保育行動為教育建議，不表示 WWF 授權或合作；未承諾任何捐款。未來衡量故事閱讀、來源開啟與回訪，尚無實績。\n'+'\n'.join(k+' '+v[0]+' '+v[1] for k,v in WILDLIFE_SOURCES.items()))

sustainability=json.loads((OUT/'sustainability_slides.json').read_text(encoding='utf-8'))
for item in sustainability['slides']:
 s=slide(item['title'],kicker=item['kicker'],sub=item['sub'])
 cards(s,item['cards'],y=2.65,h=3.55)
 text(s,item['footer'],.8,6.52,11.65,.4,12,MUTED)
 finish(s,item['short_title'],item['notes']+'\n研究與來源：docs/sustainability-direction.md；2026-09-18。')

s=slide('STEPN 的教訓：獎勵需求不能只靠再投資',kicker='經濟設計｜機制分析，不是價格預測',sub='2022 年官方描述：GST 可增發；升級與鑄鞋消耗代幣；新玩家與擴充裝備帶來購買需求。')
cards(s,[('需求的脆弱點','若消耗是為了未來賺更多\n收益预期轉弱時\n再投資需求也可能減少'),('已有的不同','免費初始鞋與免費升級\nNFT 作為成就收藏\n探索冊外觀無收益加成'),('仍未解決','固定供給仍可能領完\n單人限額不是全站預算\n競技場不是外部收入')])
text(s,'來源：STEPN 官方 2022-05-10；市場衝擊另見講稿。上述因果為本案分析。',.76,6.63,11.7,.3,11,MUTED)
finish(s,'STEPN 對照','歷史機制：https://stepnofficial.medium.com/tokeonomics-at-stepn-ee08604e82f1 。負向循環為本案機制推論，非量化因果歸因。當年另有大盤下跌與中國 GPS 限制衝擊：https://www.coindesk.com/markets/2022/05/29/first-mover-asia-bitcoin-extends-losing-streak-new-lunas-crash-like-old-lunas-stepns-china-dilemma 。不將價格下跌等同 STEPN 停止營運。完整對照 docs/economics/stepn-risk-review.md。')

s=slide('成長越快，越需要全站獎勵預算',kicker='壓力試算｜不是實際鏈上餘額',sub='假設初始金庫 200,000 tSKR；滿階每錢包每天雙任務領 33；無補庫、無回流。')
for i,(num,label) in enumerate([('60.61 天','100 個合格錢包'),('6.06 天','1,000 個合格錢包'),('0.61 天','10,000 個合格錢包')]):
 x=.74+i*4.03; rect(s,x,2.85,3.8,2.5);text(s,num,x+.2,3.28,3.4,.8,34,AMBER,True);text(s,label,x+.2,4.42,3.4,.5,18)
text(s,'公式：200,000 ÷（合格錢包數 × 33）。錢包不等於真人。',.8,5.63,11.6,.4,17)
text(s,'固定等級庫存試算；不含成長過程、行為、價格或流動性預測。',.8,6.15,11.6,.4,14,MUTED)
finish(s,'金庫壓力試算','來源：tools/tokenomics/stress.mjs 與 docs/economics/stress.csv；常數來自 Rust，未讀當前 RPC。9 組包含 100／1000／10000 錢包與 Lv1／Lv5／40 日上限。此頁選滿階不含 streak 情境。庫存／日需求為理論比值，按筆領取可能先遇不足，不是保證服務天數。固定供給限制增發，但不保證流通價格或現金價值。')

s=slide('把遊戲成長、獎勵預算與真實收入分開',kicker='經濟保護｜待實作與驗證',sub='驗收目標：沒有新玩家買幣、沒有幣價上漲，遊戲仍能提供核心體驗。')
cards(s,[('成長不中斷','打卡與代幣轉帳拆分\n金庫為零仍累積進度\n不產生未撥款欠款'),('發放有上限','全站期間預算先撥款\n版本、個人封頂與對帳\n反作弊及撤銷窗口'),('獎品有來源','依實收服務／贊助收入\n或明列的有限補貼\n先保留已承諾履約成本')])
text(s,'驗證：無獎金留存、10 倍申領、收入歸零。不能承諾幣價不跌。',.76,6.63,11.7,.3,12,AMBER)
finish(s,'經濟保護路線','PG-EC-02～05 為 TODO，未改現行合約。現況 clock_in 的代幣轉帳失敗會讓 XP 更新一起回滾；不能單靠把獎勵設為零修復。全站日預算 1000、保留 20000／180 天只是報告範例，非定案；真實獎品需實收且已撥款的成本準備。收入不含參賽本金或自家幣升值；保留既有承諾，僅調整未來活動。募資下一阶段需完成這些技術與需求驗收。')

s=slide('目前是原型階段，募資用來取得市場證據',kicker='現況與驗證計畫',sub='依現有專案文件與資產整理；尚未提供可核驗的用戶、營收或合作實績。')
cards(s,[('已具備','需求與產品設計文件\n品牌視覺與揭曉動畫\n保育跑鞋與成長盲盒外觀'),('下一步驗證','真實任務與活動端到端流程\n主辦方願付價格\n現場交付成本與錯誤率'),('尚待補齊','正式團隊履歷與分工\n已簽合作／意向書\n用戶留存與實收營收')])
finish(s,'現況','不把 PG TODO 或設計圖當成功能上線，不把 BRD 的 500 用戶、30% 留存當成實績。程式存在只能說明原型工作進行中，未在此做完整系統驗收。tSKR 為 devnet 測試代幣，無金錢價值、非官方 SKR。')

s=slide('12 個月，依驗證結果逐階段投入',kicker='募資後里程碑提案',sub='這是商業驗證排程，不取代原黑客松開發時程；各階段目標均未達成。')
cards(s,[('0–3 個月','招募 3 家設計合作夥伴\n至少完成 1 場試辦\n核驗報到、領取與成績對帳'),('4–6 個月','累計 5 場付費活動\n至少 2 家主辦方再採購\n分開追蹤軟體與現場成本'),('7–12 個月','累計 10 家付費主辦方\n建立下一期續約管線\n單場貢獻毛利轉正')])
finish(s,'里程碑','全部為建議門檻，不是承諾或業績。再採購指獨立的新付費活動；貢獻毛利＝服務收入減可歸屬雲端、支付、客服、現場與外包交付成本，不等於公司淨利。門檻未達應調整產品、客群或支出。')

s=slide('以可追溯假設，建立收入情境',kicker='第 3 年年度營收情境｜非 ARR、非財測承諾',sub='台幣萬元；服務收入不包含獎品、代收款與代幣。')
rows=[('保守','10 × 3 × 3 萬 = 90 萬','5 × 10 萬 = 50 萬','140 萬'),('基準','20 × 4 × 3 萬 = 240 萬','15 × 10 萬 = 150 萬','390 萬'),('擴張','50 × 6 × 3 萬 = 900 萬','30 × 10 萬 = 300 萬','1,200 萬')]
for i,(label,a,b,total) in enumerate(rows):
 y=2.8+i*1.06;rect(s,.75,y,11.8,.9);text(s,label,.97,y+.18,1.2,.4,20,MINT,True);text(s,a,2.28,y+.18,4,.4,17);text(s,b,6.45,y+.18,3.3,.4,17,MUTED);text(s,total,10.4,y+.14,1.9,.5,23,WHITE,True)
text(s,'主辦方 × 年場次 × 單價',2.28,2.35,4,.3,12,MUTED);text(s,'獨立品牌專案 × 單價',6.45,2.35,3.3,.3,12,MUTED);text(s,'合計',10.4,2.35,1.5,.3,12,MUTED)
text(s,'單場假設：收入 3 萬 − 直接成本 1.2 萬 = 貢獻 1.8 萬（60%）；成本待試辦驗證。',.78,6.23,11.8,.4,15,AMBER)
finish(s,'營收情境','第 3 年年化場次情境，不是 12 個月目標或 ARR。各情境公式如投影片。品牌案與活動服務的工作範圍分開計費，避免雙算。同一年度總收入保守 140 萬、基準 390 萬、擴張 1,200 萬。未含稅，未扣團隊固定成本及獲客費，無法推論損益兩平。基準情境不足支撐每年 1,200 萬支出，需進一步驗證更高客戶密度／訂閱或降低支出。')

s=slide('建議募資 NT$1,200 萬，取得 12 個月驗證窗口',kicker='募資方案｜可調整假設，尚未定案',sub='以約 NT$100 萬／月的含緩衝現金配置示意；估值、股權比例與工具待議。')
for i,(label,percent,amount,col) in enumerate([('产品與體驗',45,540,MINT),('合作試辦與客戶開發',25,300,CYAN),('營運、資料保護與法務',15,180,VIOLET),('現金緩衝',15,180,AMBER)]):
 y=2.8+i*.75;text(s,label,.8,y,4,.4,18);rect(s,4.8,y+.08,percent/10,.19,col);text(s,str(percent)+'%  /  '+str(amount)+' 萬',10,y,2.5,.4,18)
text(s,'投入門檻：付費試辦 → 再採購 → 正向單場貢獻 → 擴大銷售',.8,6.2,11.7,.45,20,MINT,True)
finish(s,'募資用途','使用者未提供金額，故採 NT$1,200 萬／12 月示意而非正式募資條件。配置 540+300+180+180=1,200 萬。並非經薪資、聘僱與現場預算驗證的 burn rate；正式募資前需依現有人力、現金、薪資及試辦報價重算。無投資報酬保證、不含代幣募資。')

s=slide('需要能同時理解運動現場與數位產品的團隊',kicker='團隊與投資人合作｜待補具名資料')
cards(s,[('產品與技術','創辦人姓名／經歷：待補\n產品與工程負責人：待補\n可展示成果與投入比例：待補'),('活動與商務','運動活動／品牌經歷：待補\n已確認合作與推薦人：待補\n客戶開發及現場交付負責人'),('期待投資人協助','跑團、賽事與品牌引介\n商業定價與通路驗證\n早期組織與募資治理')])
finish(s,'團隊與合作邀請','不虛構創辦人背景、合作方 Logo 或既有投資人。寄出正式募資版本前，必須補姓名、職稱、過往實績、股權／全職投入、聯絡資訊。')

s=slide('附錄｜市場資料與口徑',kicker='可點擊來源｜查核日 2026.09.14')
for i,(key,(label,url)) in enumerate(SOURCES.items()):
 y=2.1+i*1.38;text(s,key+'  '+label,.78,y,11.7,.5,18,MINT,True,url=url)
 desc={'S1':'2025 年發布，2024 年實績；PDF 身體活動章節 p.61–65。市場及成長數據。','S2':'2025-08-11；可比國家的平均營收成長 8%，非全球加權市場成長。','S3':'2026-07-14；逾 10 萬名用戶具 SKR 領取資格，不等於 DAU 或健身市場。'}[key]
 text(s,desc,.8,y+.62,11.5,.55,15,MUTED)
finish(s,'資料來源','\n'.join(k+' '+v[0]+'\n'+v[1] for k,v in SOURCES.items())+'\n所有外部数据查核日期 2026-09-14；本案價格、家數、收入與募資數字皆為自建情境。')


s=slide('附錄｜經濟評估的來源與重現方式',kicker='新增來源｜分析完成 2026.09.15，文件同步 2026.09.16')
items=[('E1｜STEPN 官方，2022-05-10','當時的 GST 供給、消耗與買賣動機。','https://stepnofficial.medium.com/tokeonomics-at-stepn-ee08604e82f1'),('E2｜CoinDesk，2022-05-29','大盤與中國服務限制的背景；不作單一因果歸因。','https://www.coindesk.com/markets/2022/05/29/first-mover-asia-bitcoin-extends-losing-streak-new-lunas-crash-like-old-lunas-stepns-china-dilemma'),('E3｜NeonShift 專案試算','node tools/tokenomics/stress.mjs；九組庫存情境。',None)]
for i,(label,desc,url) in enumerate(items):
 y=2.35+i*1.25;text(s,label,.8,y,11.5,.45,21,MINT,True,url=url);text(s,desc,.8,y+.53,11.5,.5,17,MUTED)
finish(s,'經濟資料來源','E1／E2 網頁於前次分析讀取，2026-09-16 沿用歷史資料；本次未重新查市場即時數據。E3: docs/economics/stepn-risk-review.md、stress.csv、programs/neonshift-core/src/constants.rs、instructions/clock_in.rs、tournament_math.rs。已驗證预算與守恆斷言；未建模價格／流動性／留存。市場來源 S1～S3 沿用 2026-09-14 查核。')

s=slide('附錄｜SOL 永續與保育研究來源',kicker='第一手來源｜查閱日期 2026.09.18',sub='技術與框架事實有來源；客戶需求、定價、任務與成效門檻均為本案提案。')
for i,(key,(label,url)) in enumerate(sustainability['sources'].items()):
 text(s,key+'  '+label,.8,2.65+i*.46,11.6,.35,15,MINT,url=url)
text(s,'C1 為 2023 年歷史報告；C0 動態值未讀取。舊市場／物種來源保留原查核日期。',.8,6.38,11.6,.45,12,MUTED)
finish(s,'永續與保育資料來源','\n'.join(k+' '+v[0]+'\n'+v[1] for k,v in sustainability['sources'].items())+'\n完整研究、假設與限制：docs/sustainability-direction.md。新增來源查閱於 2026-09-18；不代表機構合作或認證。')

s=slide('附錄｜正式對外前，補齊五項證據',kicker='募資準備清單')
items=['01  團隊姓名、專長、全職投入與股權結構','02  主辦方访談、試辦意向與可聯絡推薦人','03  示範流程、實機驗收與使用者留存資料','04  定價測試、單場成本與獲客成本','05  募資金額、現金預算、估值及使用條件']
for i,t in enumerate(items): text(s,t,.85,2.45+i*.69,11.6,.5,21,MINT if i==0 else WHITE)
finish(s,'對外前準備','本簡報為中文募資討論草稿，不代表活動合作、收入、投資條件已經成立。商業模型及募資金額以使用者補充資料後修訂。投資重點是參與經營與 B2B 合作；tSKR 是測試代幣且無金錢價值，不是官方 SKR。')

for s in prs.slides:
 for sh in s.shapes:
  assert sh.left>=0 and sh.top>=0
  assert sh.left+sh.width<=prs.slide_width+10
  assert sh.top+sh.height<=prs.slide_height+10
prs.core_properties.title='NeonShift 中文募資簡報｜討論草稿'
prs.core_properties.subject='為自己而動，為棲地同行｜Solana、永續營運與保育合作'
prs.core_properties.author='NeonShift'
target=OUT/'NeonShift_募資簡報_中文草稿_v3.pptx'; prs.save(str(target))
css='body{margin:0;background:#151925;font-family:"Microsoft JhengHei","PingFang TC",sans-serif}section{width:13.333in;height:7.5in;position:relative;background:#050711;margin:24px auto;overflow:hidden;box-shadow:0 4px 30px #000}section>*{position:absolute;box-sizing:border-box}.text{line-height:1.35;white-space:normal;overflow:hidden}.box{border-radius:12px}img{object-fit:fill}@media print{body{background:white}section{margin:0;page-break-after:always;box-shadow:none}@page{size:13.333in 7.5in;margin:0}}'
(OUT/'preview.html').write_text('<!doctype html><html lang="zh-Hant"><meta charset="utf-8"><title>NeonShift 中文募資草稿</title><style>'+css+'</style>'+''.join(html_slides)+'</html>',encoding='utf-8')
(OUT/'講稿與資料來源.md').write_text('# NeonShift 中文募資簡報講稿\n\n2026-09-18。%d 頁，%d 頁主文＋4 頁附錄。可編輯文字／圖形；視覺稿為 PNG。\n\n新增永續研究見 [共同方向與研究](../../docs/sustainability-direction.md)。\n\n'%(len(prs.slides),len(prs.slides)-4)+ '\n\n'.join('## %02d｜%s\n\n%s'%(r['slide'],r['title'],r['notes']) for r in outline),encoding='utf-8')
(OUT/'sources.json').write_text(json.dumps(dict(SOURCES, **WILDLIFE_SOURCES, **sustainability['sources'], E1=('STEPN 官方，2022-05-10', 'https://stepnofficial.medium.com/tokeonomics-at-stepn-ee08604e82f1'), E2=('CoinDesk，2022-05-29', 'https://www.coindesk.com/markets/2022/05/29/first-mover-asia-bitcoin-extends-losing-streak-new-lunas-crash-like-old-lunas-stepns-china-dilemma'), E3=('NeonShift 九組庫存試算', 'docs/economics/stress.csv')),ensure_ascii=False,indent=2),encoding='utf-8')
print('Saved %s slides: %s'%(len(prs.slides),target))
